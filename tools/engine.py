"""Layout engine for the accounting workbooks.

A question is written once as a spec function. The same spec is drawn three ways:
  mode "q": question sheet. Answer cells are yellow inputs with a check cell.
  mode "a": answer sheet. Answer cells hold formulas (green) linked to the question data.
  mode "w": worked example. Answer cells hold formulas; inputs are blue and editable.
Because every mode draws the same rows, a check cell on the question sheet compares
with the cell at the same address on the answer sheet.

Every block is also recorded (self.blocks) so tools/build_docs.py can render the
same content as Markdown for the notes, model questions and answer key.
"""
import math
import re

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ARIAL = "Arial"
NAVY = "1F3864"
BLUE = "0000FF"
GREEN = "008000"
GREY = "595959"


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


F_TITLE = fill(NAVY)
F_HEAD = fill("D9E1F2")
F_PART = fill("DDEBF7")
F_INPUT = fill("FFF59D")
F_ANS = fill("E2EFDA")
F_DATA = fill("F2F2F2")
F_WRITE = fill("EEF6E8")

NUM = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = "0.0%"
INT = "0"
TXT = "General"

LIGHT = Side(style="thin", color="B4C6E7")
DARK = Side(style="thin", color="000000")
DOUBLE = Side(style="double", color="000000")
MED = Side(style="medium", color=NAVY)
NONE = Side()

U = 31          # grid width in units
UW = 4.6        # width of one unit column
C0 = 2          # first grid column (B)
TICK, CROSS = "✓", "✗"
STAT_COLS = {1: 16, 2: 21, 3: 26}   # unit offsets of the three RM columns


def font(bold=False, italic=False, color="000000", size=10):
    return Font(name=ARIAL, size=size, bold=bold, italic=italic, color=color)


def edge(line=None, boxed=False):
    side = LIGHT if boxed else NONE
    top = DARK if line in ("top", "total") else side
    bottom = DOUBLE if line == "total" else DARK if line == "under" else side
    return Border(left=side, right=side, top=top, bottom=bottom)


YEAR = "@Y"


def split_year(date):
    """'2025 Mar 31' -> ('2025', 'Mar 31'); 'Sep 30' -> (None, 'Sep 30')."""
    m = re.match(r"^(\d{4}(?:-\d{2})?|Year \d+)\b\s*(.*)$", date or "")
    return (m.group(1), m.group(2)) if m else (None, date or "")


def expand_journal(rows):
    """Put the year on its own row (with RM on the first one) and keep only month and day on entries."""
    out, cur, first = [], None, True
    for row in rows:
        if row[0] == "n":
            out.append(row)
            continue
        kind, date, text, t = row
        y, rest = split_year(date)
        if first and y is None:
            out.append(("y", "", True))
            first = False
        if y and y != cur:
            out.append(("y", y, first))
            first, cur = False, y
        out.append((kind, rest, text, t))
    return out


def expand_ledger(rows):
    """Same idea for a ledger: each side gets a year row (RM on the first) whenever the year changes."""
    from itertools import zip_longest
    sections, cur = [], {"dr": [], "cr": []}
    for row in rows:
        if row == "TOTAL":
            sections.append((cur, True))
            cur = {"dr": [], "cr": []}
            continue
        for side, e in zip(("dr", "cr"), row):
            if e:
                cur[side].append(e)
    if cur["dr"] or cur["cr"]:
        sections.append((cur, False))
    state, first, out = {"dr": None, "cr": None}, {"dr": True, "cr": True}, []
    for cur, total in sections:
        lists = {}
        for side in ("dr", "cr"):
            items = []
            for e in cur[side]:
                y, rest = split_year(e[0])
                if first[side] and y is None:
                    items.append((YEAR, "", True))
                    first[side] = False
                if y and y != state[side]:
                    items.append((YEAR, y, first[side]))
                    first[side], state[side] = False, y
                items.append((rest,) + tuple(e[1:]))
            lists[side] = items
        out += list(zip_longest(lists["dr"], lists["cr"]))
        if total:
            out.append("TOTAL")
    return out



class V:
    """An answer value: a formula template with optional underline style and key."""

    def __init__(self, t, line=None, key=None, fmt=None):
        self.t, self.line, self.key, self.fmt = str(t), line, key, fmt


def D(value, key, fmt=None):
    """A data (input) cell inside a table."""
    return ("in", value, key, fmt)


def T(text):
    """A plain text cell inside a table."""
    return ("t", text)


def lines_needed(text, span):
    per = max(8, int(span * UW * 1.18))
    return sum(max(1, math.ceil(len(p) / per)) for p in str(text).split("\n"))


def as_v(spec):
    return spec if isinstance(spec, V) else V(spec)


class Page:
    def __init__(self, wb, name, mode, twin=None, tab=None):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 110
        ws.column_dimensions["A"].width = 2
        for i in range(U + 2):
            ws.column_dimensions[get_column_letter(C0 + i)].width = UW
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        if tab:
            ws.sheet_properties.tabColor = tab
        self.ws, self.name, self.mode, self.twin = ws, name, mode, twin
        self.r = 1
        self.ids = {}
        self.todo = []
        self.inputs = 0
        self.blocks = []
        self.score = None
        self.first_row = 1
        self._tag = None

    # ------------------------------------------------------------------ basics
    def addr(self, r, u, absolute=False):
        c = get_column_letter(C0 + u)
        return f"${c}${r}" if absolute else f"{c}{r}"

    def box(self, r, u, span, value=None, bold=False, italic=False, color="000000", fill_=None,
            align="left", wrap=False, fmt=None, border=None, size=10, rows=1, indent=0, valign=None):
        ws = self.ws
        c = C0 + u
        if span > 1 or rows > 1:
            ws.merge_cells(start_row=r, start_column=c, end_row=r + rows - 1, end_column=c + span - 1)
        for rr in range(r, r + rows):
            for cc in range(c, c + span):
                x = ws.cell(row=rr, column=cc)
                if fill_:
                    x.fill = fill_
                if border:
                    x.border = border
        cell = ws.cell(row=r, column=c)
        if value is not None:
            cell.value = value
        cell.font = font(bold, italic, color, size)
        cell.alignment = Alignment(horizontal=align, wrap_text=wrap, indent=indent,
                                   vertical=valign or ("top" if (wrap or rows > 1) else "center"))
        if fmt:
            cell.number_format = fmt
        return cell

    def height(self, r, h):
        cur = self.ws.row_dimensions[r].height or 0
        self.ws.row_dimensions[r].height = max(cur, h)

    def tag(self, name):
        self._tag = name
        return self

    def _rec(self, block):
        if self._tag:
            block["tag"] = self._tag
            self._tag = None
        self.blocks.append(block)

    def val(self, r, u, span, spec, fmt=NUM, check=True, grid=False):
        """Draw an answer cell. Returns its relative address."""
        spec = as_v(spec)
        fmt = spec.fmt or fmt
        a = self.addr(r, u)
        if spec.key:
            self.ids[spec.key] = self.addr(r, u, True)
        align = "center" if fmt == TXT else "right"
        if self.mode == "q":
            self.box(r, u, span, fill_=F_INPUT, fmt=fmt, border=edge(spec.line, True), align=align)
            self.inputs += 1
            if check:
                self.box(r, u + span, 1,
                         f'=IF({a}="","",IFERROR(IF(ABS(ABS({a})-ABS(\'{self.twin}\'!{a}))<0.5,'
                         f'"{TICK}","{CROSS}"),"{CROSS}"))',
                         bold=True, align="center", border=edge(None, grid))
        else:
            cell = self.box(r, u, span, fill_=F_ANS if self.mode == "a" else None, bold=self.mode == "a",
                            fmt=fmt, border=edge(spec.line, grid or self.mode == "a"), align=align)
            self.todo.append((cell, spec.t))
            if check and grid:
                self.box(r, u + span, 1, border=edge(None, True))
        return a

    def datum(self, r, u, span, value, key, fmt=NUM):
        a = self.addr(r, u)
        self.ids[key] = self.addr(r, u, True)
        fmt = fmt or NUM
        if self.mode == "a":
            self.box(r, u, span, f"='{self.twin}'!{a}", color=GREEN, fmt=fmt, align="right",
                     fill_=F_DATA, border=edge(None, True))
        else:
            self.box(r, u, span, value, color=BLUE, fmt=fmt, align="right", fill_=F_DATA,
                     border=edge(None, True))
        return a

    def finish(self):
        for cell, t in self.todo:
            f = re.sub(r"\{(\w+)\}", lambda m: self.ids[m.group(1)], t)
            cell.value = f if f.startswith("=") else "=" + f
        last = get_column_letter(C0 + U)
        end = max(self.r, 2)
        body = f"B{self.first_row}:{last}{end}"
        if self.score:
            r, u = self.score
            self.ws.cell(row=r, column=C0 + u).value = (
                f'=COUNTIF({body},"{TICK}")&" correct out of {self.inputs} answer cells"')
        self.ws.conditional_formatting.add(body, CellIsRule(
            operator="equal", formula=[f'"{TICK}"'], font=Font(name=ARIAL, color="00B050", bold=True)))
        self.ws.conditional_formatting.add(body, CellIsRule(
            operator="equal", formula=[f'"{CROSS}"'], font=Font(name=ARIAL, color="C00000", bold=True)))
        self.ws.print_area = f"A1:{last}{end}"

    # ------------------------------------------------------------------ text blocks
    def title(self, text, sub=None):
        r = self.r
        self.box(r, 0, U, text, bold=True, color="FFFFFF", fill_=F_TITLE, size=13, indent=1)
        self.height(r, 26)
        self.r += 1
        self._rec({"k": "title", "text": text})
        if sub:
            self.note(sub)

    def status(self):
        """Two instruction rows plus a spacer; identical row count in every mode."""
        r = self.r
        if self.mode == "q":
            self.box(r, 0, U, "Type your answers in the yellow cells. Numbers only: no RM, no commas. "
                     "Enter 0 where the answer is nil. ✓ = correct, ✗ = try again.",
                     italic=True, color=GREY)
            self.box(r + 1, 0, 6, "Your score:", bold=True, color=NAVY)
            self.box(r + 1, 6, 25, bold=True, color=NAVY)
            self.score = (r + 1, 6)
            self.first_row = r + 2
        elif self.mode == "a":
            self.box(r, 0, U, "Answer sheet. Green cells hold the correct answers. Each figure is a formula "
                     "linked to the question sheet, so the answers follow any change to the blue figures.",
                     italic=True, color=GREY)
            self.box(r + 1, 0, U, "Written answers are model answers: other wording that makes the same "
                     "points also earns the marks.", italic=True, color=GREY)
        else:
            self.box(r, 0, U, "Worked example. Change any blue figure and every answer below updates.",
                     italic=True, color=GREY)
            self.box(r + 1, 0, U, "Black figures are formulas. Do not type over them.", italic=True, color=GREY)
        self.r += 3

    def text(self, t, bold=False, italic=False, color="000000", indent=0, record=True, style="p"):
        r = self.r
        span = U - indent
        self.box(r, indent, span, t, bold=bold, italic=italic, color=color, wrap=True)
        self.height(r, 13.5 * lines_needed(t, span) + 3)
        self.r += 1
        if record:
            self._rec({"k": "text", "text": t, "style": style})

    def bullets(self, items):
        for it in items:
            self.text("•  " + it, indent=1, record=False)
            self._rec({"k": "text", "text": it, "style": "bullet"})

    def note(self, t):
        self.text(t, italic=True, color=GREY, record=False)

    def gap(self, n=1):
        self.r += n

    def part(self, t):
        self.gap()
        r = self.r
        self.box(r, 0, U, t, bold=True, color=NAVY, fill_=F_PART, wrap=True, border=Border(left=MED))
        self.height(r, 14 * lines_needed(t, U) + 5)
        self.r += 1
        self._rec({"k": "part", "text": t})

    # ------------------------------------------------------------------ data blocks
    def data(self, rows, title="Figures from the question", md=False):
        r = self.r
        self.box(r, 0, 22, title, bold=True, fill_=F_HEAD, border=edge(None, True))
        self.box(r, 22, 6, "Figure", bold=True, fill_=F_HEAD, align="center", border=edge(None, True))
        self.r += 1
        rec = []
        for row in rows:
            if row[0] is None:
                self.box(self.r, 0, 28, row[1], bold=True, italic=True, border=edge(None, True))
                rec.append(("#", row[1], None))
                self.r += 1
                continue
            key, label, value = row[:3]
            fmt = row[3] if len(row) > 3 else NUM
            self.box(self.r, 0, 22, label, border=edge(None, True), wrap=True)
            self.height(self.r, 13.5 * lines_needed(label, 22) + 3)
            self.datum(self.r, 22, 6, value, key, fmt)
            rec.append((label, value, fmt))
            self.r += 1
        self._rec({"k": "data", "title": title, "md": md, "rows": rec})
        self.gap()

    def tb(self, rows, title="Trial balance"):
        r = self.r
        self.box(r, 0, U, title, bold=True, color=NAVY, align="center")
        r += 1
        for u, span, h in ((0, 19, "Particulars"), (19, 6, "Debit (RM)"), (25, 6, "Credit (RM)")):
            self.box(r, u, span, h, bold=True, fill_=F_HEAD, align="left" if u == 0 else "center",
                     border=edge(None, True))
        r += 1
        first = r
        rec = []
        for key, label, dr, cr in rows:
            self.box(r, 0, 19, label, border=edge(None, True))
            for u, v in ((19, dr), (25, cr)):
                if v is None:
                    self.box(r, u, 6, border=edge(None, True))
                else:
                    self.datum(r, u, 6, v, key)
            rec.append((label, dr, cr))
            r += 1
        self.box(r, 0, 19, "Total", bold=True, border=edge(None, True))
        for u in (19, 25):
            cell = self.box(r, u, 6, bold=True, fmt=NUM, align="right", border=edge("total", True))
            self.todo.append((cell, f"SUM({self.addr(first, u)}:{self.addr(r - 1, u)})"))
        self.r = r + 1
        dsum = sum(x[2] for x in rows if x[2] is not None)
        csum = sum(x[3] for x in rows if x[3] is not None)
        self._rec({"k": "tb", "title": title, "rows": rec, "totals": (dsum, csum)})
        self.gap()

    # ------------------------------------------------------------------ answer blocks
    def calc(self, rows):
        rec = []
        for row in rows:
            key, label, t = row[:3]
            fmt = row[3] if len(row) > 3 else NUM
            self.box(self.r, 0, 22, label, wrap=True, valign="center", border=Border(bottom=LIGHT))
            self.height(self.r, max(17, 13.5 * lines_needed(label, 22) + 4))
            a = self.val(self.r, 22, 5, V(t, key=key, fmt=fmt))
            rec.append((label, a, fmt))
            self.r += 1
        self._rec({"k": "calc", "rows": rec})

    def choice(self, rows):
        rec = []
        for row in rows:
            key, label, answer, options = row[:4]
            accept = row[4] if len(row) > 4 else None
            r = self.r
            self.box(r, 0, 18, label, wrap=True, valign="center", border=Border(bottom=LIGHT))
            self.height(r, max(17, 13.5 * lines_needed(label, 18) + 4))
            a = self.addr(r, 18)
            self.ids[key] = self.addr(r, 18, True)
            if self.mode == "q":
                self.box(r, 18, 9, fill_=F_INPUT, border=edge(None, True), align="center")
                dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
                self.ws.add_data_validation(dv)
                dv.add(a)
                self.inputs += 1
                if accept:
                    cond = ",".join(f'LOWER(TRIM({a}))="{x.lower()}"' for x in accept)
                    test = f"OR({cond})"
                else:
                    test = f"LOWER(TRIM({a}))=LOWER('{self.twin}'!{a})"
                self.box(r, 27, 1, f'=IF({a}="","",IF({test},"{TICK}","{CROSS}"))', bold=True, align="center")
            else:
                cell = self.box(r, 18, 9, fill_=F_ANS if self.mode == "a" else None, bold=True,
                                border=edge(None, True), align="center", wrap=True)
                self.todo.append((cell, answer))
            rec.append((label, a, accept))
            self.r += 1
        self._rec({"k": "choice", "rows": rec})

    def written(self, model, lines=3):
        r = self.r
        need = 13.5 * lines_needed(model, U) + 8
        per = max(17, need / lines)
        if self.mode == "q":
            self.box(r, 0, U, fill_=F_INPUT, border=edge(None, True), rows=lines, wrap=True)
        else:
            self.box(r, 0, U, model, fill_=F_WRITE, border=edge(None, True), rows=lines, wrap=True)
        for i in range(lines):
            self.height(r + i, per)
        self.r += lines
        self._rec({"k": "written", "text": model})

    def table(self, headers, rows, label_span, col_span, checks=True, fmts=None, total_rows=()):
        step = col_span + (1 if checks else 0)
        r = self.r
        self.box(r, 0, label_span, headers[0], bold=True, fill_=F_HEAD, border=edge(None, True), wrap=True)
        h = max(lines_needed(x, step) for x in headers[1:])
        for i, head in enumerate(headers[1:]):
            self.box(r, label_span + i * step, step, head, bold=True, fill_=F_HEAD, align="center",
                     wrap=True, border=edge(None, True))
        self.height(r, 13.5 * max(h, lines_needed(headers[0], label_span)) + 4)
        self.r += 1
        rec = []
        for label, cells in rows:
            r = self.r
            self.box(r, 0, label_span, label, border=edge(None, True), wrap=True,
                     bold=label in total_rows)
            self.height(r, max(17, 13.5 * lines_needed(label, label_span) + 4))
            out = []
            for i, c in enumerate(cells):
                u = label_span + i * step
                fmt = (fmts[i] if fmts else None) or NUM
                if c is None:
                    self.box(r, u, step, border=edge(None, True))
                    out.append(None)
                elif isinstance(c, tuple) and c[0] == "in":
                    self.datum(r, u, col_span, c[1], c[2], c[3] or fmt)
                    if checks:
                        self.box(r, u + col_span, 1, border=edge(None, True))
                    out.append(("d", c[1], c[3] or fmt))
                elif isinstance(c, tuple) and c[0] == "t":
                    self.box(r, u, step, c[1], border=edge(None, True), align="center", wrap=True)
                    out.append(("t", c[1]))
                else:
                    a = self.val(r, u, col_span, c, fmt, check=checks, grid=True)
                    out.append(("v", a, as_v(c).fmt or fmt))
            rec.append((label, out))
            self.r += 1
        self._rec({"k": "table", "headers": headers, "rows": rec})

    def journal(self, rows):
        r = self.r
        for u, span, h in ((0, 3, "Date"), (3, 16, "Particulars"), (19, 6, "Debit"), (25, 6, "Credit")):
            self.box(r, u, span, h, bold=True, fill_=F_HEAD, align="left" if u < 19 else "center",
                     border=edge(None, True))
        self.r += 1
        rec = []
        g = edge(None, True)
        for row in expand_journal(rows):
            r = self.r
            if row[0] == "y":
                self.box(r, 0, 3, row[1], bold=True, border=g)
                self.box(r, 3, 16, border=g)
                for u in (19, 25):
                    self.box(r, u, 6, "RM" if row[2] else None, bold=True, align="center", border=g)
                rec.append(("y", row[1], row[2]))
            elif row[0] == "n":
                self.box(r, 0, 3, border=g)
                self.box(r, 3, 16, "(" + row[1] + ")", italic=True, color=GREY, border=g, wrap=True)
                self.height(r, max(17, 13.5 * lines_needed(row[1], 16) + 4))
                self.box(r, 19, 6, border=g)
                self.box(r, 25, 6, border=g)
                rec.append(("n", row[1]))
            else:
                kind, date, text, t = row
                self.box(r, 0, 3, date, border=g)
                self.box(r, 3, 16, text, border=g, indent=3 if kind == "c" else 0)
                if kind == "d":
                    a = self.val(r, 19, 5, t, grid=True)
                    self.box(r, 25, 6, border=g)
                else:
                    self.box(r, 19, 6, border=g)
                    a = self.val(r, 25, 5, t, grid=True)
                rec.append((kind, date, text, a))
            self.height(r, 17)
            self.r += 1
        self._rec({"k": "journal", "rows": rec})

    def ledger(self, name, rows):
        r = self.r
        self.box(r, 0, 3, "Dr", bold=True, color=NAVY)
        self.box(r, 3, 25, name, bold=True, align="center", color=NAVY)
        self.box(r, 28, 3, "Cr", bold=True, align="right", color=NAVY)
        self.r += 1
        r = self.r
        heads = ((0, 3, "Date"), (3, 7, "Particulars"), (10, 1, "Folio"), (11, 4, "Amount"))
        for o in (0, 16):
            for u, span, h in heads:
                self.box(r, o + u, span, h, bold=True, fill_=F_HEAD, align="center" if u >= 10 else "left",
                         border=Border(bottom=MED), size=9 if u == 10 else 10)
        self.box(r, 15, 1, border=Border(bottom=MED))
        self.ws.cell(row=r, column=C0 + 16).border = Border(left=MED, bottom=MED)
        self.r += 1
        start = self.r
        rec = []
        for row in expand_ledger(rows):
            r = self.r
            if row == "TOTAL":
                entry = {"total": True}
                for o, side in ((0, "dr"), (16, "cr")):
                    rng = f"{self.addr(start, o + 11)}:{self.addr(r - 1, o + 11)}"
                    entry[side] = self.val(r, o + 11, 3, V(f"SUM({rng})", line="total"))
                rec.append(entry)
                start = r + 1
            else:
                entry = {}
                for o, side, e in ((0, "dr", row[0]), (16, "cr", row[1])):
                    if e and e[0] == YEAR:
                        self.box(r, o, 3, e[1], bold=True)
                        if e[2]:
                            self.box(r, o + 11, 3, "RM", bold=True, align="right")
                        entry[side] = [YEAR, e[1], e[2]]
                    elif e:
                        date, part, folio, t = e
                        self.box(r, o, 3, date)
                        self.box(r, o + 3, 7, part, wrap=True, valign="center")
                        self.height(r, 13.5 * lines_needed(part, 7) + 4)
                        self.box(r, o + 10, 1, folio, align="center", color=GREY, size=9)
                        entry[side] = (date, part, folio, self.val(r, o + 11, 3, t))
                    else:
                        entry[side] = None
                rec.append(entry)
            c = self.ws.cell(row=r, column=C0 + 16)
            c.border = Border(left=MED, top=c.border.top, bottom=c.border.bottom)
            self.height(r, 17)
            self.r += 1
        self._rec({"k": "ledger", "name": name, "rows": rec})

    def statement(self, heading, rows, kind="spl"):
        for h in heading:
            self.box(self.r, 0, U, h, bold=True, color=NAVY, align="center")
            self.r += 1
        r = self.r
        self.box(r, 0, 3, fill_=F_HEAD, border=edge(None, True))
        self.box(r, 3, 13, "Particulars", bold=True, fill_=F_HEAD, border=edge(None, True))
        for c, u in STAT_COLS.items():
            self.box(r, u, 5, "RM", bold=True, fill_=F_HEAD, align="center", border=edge(None, True))
        self.r += 1
        if kind == "sofp":
            r = self.r
            self.box(r, 0, 3, fill_=F_HEAD, border=edge(None, True))
            self.box(r, 3, 13, fill_=F_HEAD, border=edge(None, True))
            for (c, u), h in zip(STAT_COLS.items(), ("Cost", "Accumulated Depreciation", "Carrying Amount")):
                self.box(r, u, 5, h, bold=True, fill_=F_HEAD, align="center", wrap=True, border=edge(None, True))
            self.height(r, 30)
            self.r += 1
        rec = []
        for row in rows:
            m, p, vals = row[:3]
            flags = row[3] if len(row) > 3 else ""
            r = self.r
            self.box(r, 0, 3, m, italic=True)
            self.box(r, 3, 13, p, bold=("h" in flags or "b" in flags), wrap=True, valign="center",
                     color=NAVY if "h" in flags else "000000")
            self.height(r, max(17, 13.5 * lines_needed(p, 13) + 4))
            out = {}
            for col, spec in vals.items():
                spec = as_v(spec)
                out[col] = (self.val(r, STAT_COLS[col], 4, spec), spec.line)
            rec.append((m, p, out, flags))
            self.r += 1
        self._rec({"k": "statement", "kind": kind, "heading": heading, "rows": rec})


# ---------------------------------------------------------------------- statement helpers
def _sum(keys):
    return "+".join("{" + k + "}" for k in keys)


def spl_rows(rev, oi, pur, close, incomes=(), expenses=(), sr=None, pr=None, ci=None, to_gp=False):
    """Rows for a Statement of Profit or Loss in the three-column RM layout."""
    rows = [("", "Revenue", {3: V(rev, key="s_rev")})]
    nr = "{s_rev}"
    if sr:
        rows += [("Less:", "Sales returns", {3: V(f"-({sr})", line="under", key="s_sr")}),
                 ("", "Net revenue", {3: V("{s_rev}+{s_sr}", key="s_nr")})]
        nr = "{s_nr}"
    rows += [("Less:", "Cost of sales", {}, "h"),
             ("", "Opening inventory", {2: V(oi, key="s_oi")})]
    plines = [("Add:", "Purchases", pur)]
    if pr:
        plines.append(("Less:", "Purchases returns", f"-({pr})"))
    if ci:
        plines.append(("Add:", "Carriage inwards", ci))
    if len(plines) == 1:
        rows.append(("Add:", "Purchases", {2: V(pur, line="under", key="s_np")}))
    else:
        keys = []
        for i, (m, p, t) in enumerate(plines):
            k = f"s_p{i}"
            keys.append(k)
            last = i == len(plines) - 1
            vals = {1: V(t, key=k, line="under" if last else None)}
            if last:
                vals[2] = V(_sum(keys), key="s_np", line="under")
            rows.append((m, p, vals))
    rows += [("", "Cost of goods available for sale", {2: V("{s_oi}+{s_np}", key="s_cga")}),
             ("Less:", "Closing inventory", {2: V(f"-({close})", line="under", key="s_cl"),
                                             3: V("-({s_cga}+{s_cl})", line="under", key="s_cos")}),
             ("", "Gross profit", {3: V(nr + "+{s_cos}", key="s_gp", line="total" if to_gp else None)}, "b")]
    if to_gp:
        return rows
    sub = "{s_gp}"
    if incomes:
        rows.append(("Add:", "Other income", {}, "h"))
        keys = []
        for i, (p, t) in enumerate(incomes):
            k = f"s_i{i}"
            keys.append(k)
            last = i == len(incomes) - 1
            vals = {2: V(t, key=k, line="under" if last else None)}
            if last:
                vals[3] = V(_sum(keys), key="s_it", line="under")
            rows.append(("", p, vals))
        rows.append(("", "", {3: V("{s_gp}+{s_it}", key="s_sub")}))
        sub = "{s_sub}"
    rows.append(("Less:", "Expenses", {}, "h"))
    keys = []
    for i, (p, t) in enumerate(expenses):
        k = f"s_e{i}"
        keys.append(k)
        last = i == len(expenses) - 1
        vals = {2: V(t, key=k, line="under" if last else None)}
        if last:
            vals[3] = V(f"-({_sum(keys)})", key="s_et", line="under")
        rows.append(("", p, vals))
    rows.append(("", "Profit for the year", {3: V(sub + "+{s_et}", key="s_profit", line="total")}, "b"))
    return rows


def sofp_rows(nca, ca, cl, opening, profit, drawings, ncl=()):
    """Rows for a Statement of Financial Position.

    nca: [(label, cost, accumulated depreciation)]
    ca:  [(label, amount)] or ("TR", trade receivables, allowance)
    cl, ncl: [(label, amount)]
    """
    rows = [("", "Non-current assets", {}, "h")]
    many = len(nca) > 1
    for i, (p, c, a) in enumerate(nca):
        ln = "under" if many and i == len(nca) - 1 else None
        rows.append(("", p, {1: V(c, key=f"f_c{i}", line=ln), 2: V(a, key=f"f_a{i}", line=ln),
                             3: V(f"{{f_c{i}}}-{{f_a{i}}}", key=f"f_n{i}", line=ln)}))
    if many:
        n = range(len(nca))
        rows.append(("", "Total non-current assets", {1: V(_sum(f"f_c{i}" for i in n)),
                                                     2: V(_sum(f"f_a{i}" for i in n)),
                                                     3: V(_sum(f"f_n{i}" for i in n), key="f_nca")}, "b"))
        nca_ref = "{f_nca}"
    else:
        nca_ref = "{f_n0}"
    rows.append(("", "Current assets", {}, "h"))
    keys = []
    for i, item in enumerate(ca):
        k = f"f_ca{i}"
        keys.append(k)
        ln = "under" if i == len(ca) - 1 else None
        if item[0] == "TR":
            rows.append(("", "Trade receivables", {1: V(item[1], key="f_tr")}))
            rows.append(("Less:", "Allowance for irrecoverable debts",
                         {1: V(f"-({item[2]})", key="f_al", line="under"), 2: V("{f_tr}+{f_al}", key=k, line=ln)}))
        else:
            rows.append(("", item[0], {2: V(item[1], key=k, line=ln)}))
    rows.append(("", "Total current assets", {2: V(_sum(keys), key="f_cat")}, "b"))
    rows.append(("Less:", "Current liabilities", {}, "h"))
    keys = []
    for i, (p, t) in enumerate(cl):
        k = f"f_cl{i}"
        keys.append(k)
        rows.append(("", p, {2: V(t, key=k, line="under" if i == len(cl) - 1 else None)}))
    rows.append(("", "Total current liabilities", {2: V(_sum(keys), key="f_clt", line="under")}, "b"))
    rows.append(("", "Net current assets (working capital)", {3: V("{f_cat}-{f_clt}", key="f_wc", line="under")}))
    if ncl:
        rows.append(("", "", {3: V(nca_ref + "+{f_wc}", key="f_tot")}))
        rows.append(("Less:", "Non-current liabilities", {}, "h"))
        keys = []
        for i, (p, t) in enumerate(ncl):
            k = f"f_nl{i}"
            keys.append(k)
            rows.append(("", p, {3: V(f"-({t})", key=k, line="under" if i == len(ncl) - 1 else None)}))
        rows.append(("", "Net assets", {3: V("{f_tot}+" + _sum(keys), key="f_na", line="total")}, "b"))
    else:
        rows.append(("", "Net assets", {3: V(nca_ref + "+{f_wc}", key="f_na", line="total")}, "b"))
    rows += [("", "Financed by:", {}, "h"),
             ("", "Capital at start of year", {3: V(opening, key="f_op")}),
             ("Add:", "Profit for the year", {3: V(profit, key="f_pr", line="under")}),
             ("", "", {3: V("{f_op}+{f_pr}", key="f_s2")}),
             ("Less:", "Drawings", {3: V(f"-({drawings})", key="f_dr", line="under")}),
             ("", "Capital at end of year", {3: V("{f_s2}+{f_dr}", key="f_cc", line="total")}, "b")]
    return rows
