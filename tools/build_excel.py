"""Build the Excel workbooks for the Senior 1 accounting notes.

Run from the repository root:  python3 tools/build_excel.py
Then recalculate each file (for example open and save in Excel, or LibreOffice).
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "excel"

ARIAL = "Arial"
BLUE = "0000FF"
YELLOW = PatternFill("solid", start_color="FFFF00")
HEAD = PatternFill("solid", start_color="D9E1F2")
NUM = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = "0.0%"
THIN = Side(style="thin", color="999999")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
TOP = Border(top=Side(style="thin", color="000000"))
TOTAL = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))
GREEN_FONT = Font(name=ARIAL, color="006100", bold=True)
RED_FONT = Font(name=ARIAL, color="9C0006", bold=True)


def font(**kw):
    return Font(name=ARIAL, size=kw.pop("size", 11), **kw)


def widths(ws, **cols):
    for col, w in cols.items():
        ws.column_dimensions[col].width = w


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = font(bold=True, size=14)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = font(italic=True, color="555555")


def text(ws, ref, value, bold=False, italic=False, wrap=False):
    c = ws[ref]
    c.value = value
    c.font = font(bold=bold, italic=italic)
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def header(ws, row, values, col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font = font(bold=True)
        c.fill = HEAD
        c.border = BOX
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def inp(ws, ref, value, fmt=NUM):
    c = ws[ref]
    c.value = value
    c.font = font(color=BLUE)
    c.number_format = fmt
    return c


def calc(ws, ref, formula, fmt=NUM, bold=False, border=None):
    c = ws[ref]
    c.value = formula
    c.font = font(bold=bold)
    c.number_format = "General" if fmt == "@" else fmt
    if border:
        c.border = border
    return c


def check_rules(ws, rng):
    first = rng.split(":")[0]
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first},7)="Correct"'], font=GREEN_FONT))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first},3)="Try"'], font=RED_FONT))


def read_me(wb, topic, extra_lines=()):
    ws = wb.active
    ws.title = "Read Me"
    widths(ws, A=4, B=100)
    title(ws, topic, "Senior 1 Accounting. Free to use and share.")
    lines = [
        ("How to use this workbook", True),
        ("1. Start with the worked example sheets. Change any blue number and watch every answer update.", False),
        ("2. Then try the Practice sheet. Type your answers in the yellow cells.", False),
        ("3. The Check column tells you 'Correct' or 'Try again' as soon as you type.", False),
        ("4. The Answers sheet is hidden. When you finish, right-click any sheet tab and choose Unhide to see full answers.", False),
        ("", False),
        ("Colour key", True),
        ("Blue text = an input you can change.", False),
        ("Black text = a formula. Do not type over it.", False),
        ("Yellow cell = your answer goes here.", False),
        ("", False),
    ]
    lines += [(l, False) for l in extra_lines]
    r = 4
    for line, bold in lines:
        text(ws, f"B{r}", line, bold=bold)
        r += 1
    ws["B11"].font = font(color=BLUE)
    ws["B13"].fill = YELLOW
    return ws


def practice(wb, heading, intro, build_info, questions, options_sheet=None):
    """Create a Practice sheet and a hidden Answers sheet with matching rows.

    build_info(ws, row) writes the question data and returns (refs, next_row),
    where refs maps a key to an absolute cell address on the Practice sheet.
    questions: list of str (section heading) or tuples
        (label, formula_template, kind, fmt)  kind is "num" or "text".
    Templates use {key} placeholders which become Practice!$B$7 style refs.
    For "text" questions the template is a formula giving the expected text,
    and options is the drop-down source (a quoted list or a cell range).
    """
    ws = wb.create_sheet("Practice")
    an = wb.create_sheet("Answers")
    widths(ws, A=58, B=18, C=18, D=14, E=14)
    widths(an, A=58, B=18)
    title(ws, heading, "Type your answers in the yellow cells. The Check column marks them.")
    title(an, heading + ": Answers", "Every answer is a formula linked to the Practice sheet data.")
    r = 3
    for line in intro:
        text(ws, f"A{r}", line, wrap=False)
        r += 1
    r += 1
    refs, r = build_info(ws, r)
    refs = {k: f"Practice!{v}" for k, v in refs.items()}
    r += 1
    header(ws, r, ["Question", "Your answer", "Check"])
    header(an, r, ["Question", "Answer"])
    r += 1
    first = r
    for q in questions:
        if isinstance(q, str):
            text(ws, f"A{r}", q, bold=True)
            text(an, f"A{r}", q, bold=True)
            r += 1
            continue
        label, tmpl, kind, fmt, *rest = q
        text(ws, f"A{r}", label)
        text(an, f"A{r}", label)
        cell = ws[f"B{r}"]
        cell.fill = YELLOW
        cell.border = BOX
        cell.number_format = fmt
        if kind == "num":
            formula = "=" + tmpl.format(**refs)
            calc(an, f"B{r}", formula, fmt, bold=True)
            ws[f"C{r}"] = (
                f'=IF(B{r}="","",IFERROR(IF(ABS(B{r}-Answers!B{r})<0.5,"Correct","Try again"),"Enter a number"))'
            )
        else:
            options = rest[0]
            calc(an, f"B{r}", tmpl.format(**refs), bold=True)
            dv = DataValidation(type="list", formula1=options, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"B{r}")
            ws[f"C{r}"] = f'=IF(B{r}="","",IF(LOWER(TRIM(B{r}))=LOWER(Answers!B{r}),"Correct","Try again"))'
        ws[f"C{r}"].font = font(bold=True)
        r += 1
    check_rules(ws, f"C{first}:C{r}")
    score_row = r + 1
    text(ws, f"A{score_row}", "Your score (number correct)", bold=True)
    calc(ws, f"B{score_row}", f'=COUNTIF(C{first}:C{r},"Correct")&" / "&COUNTA(B{first}:B{r})&" answered"', "@", bold=True)
    an.sheet_state = "hidden"
    return ws, an


def info_rows(ws, r, rows, col_label="A", col_val="B"):
    """rows: list of (key or None, label, value, fmt). None key = heading."""
    refs = {}
    for key, label, value, fmt in rows:
        if key is None:
            text(ws, f"{col_label}{r}", label, bold=True)
        else:
            text(ws, f"{col_label}{r}", label)
            inp(ws, f"{col_val}{r}", value, fmt)
            refs[key] = f"${col_val}${r}"
        r += 1
    return refs, r


# ---------------------------------------------------------------- Chapter 8
def build_ch08():
    wb = Workbook()
    read_me(wb, "Chapter 8: Irrecoverable Debts and Allowance for Receivables")

    ws = wb.create_sheet("Allowance Calculator")
    widths(ws, A=52, B=16, C=16, D=16)
    title(ws, "Worked example 8.3: allowance for irrecoverable debts",
          "Change the blue numbers. Everything else recalculates.")
    header(ws, 4, ["Item", "2023", "2024", "2025"])
    ws["A5"], ws["A6"] = "Trade receivables before write-offs", "Less: irrecoverable debts written off"
    for col, before, wo in (("B", 20500, 500), ("C", 26800, 800), ("D", 18600, 600)):
        inp(ws, f"{col}5", before)
        inp(ws, f"{col}6", wo)
        calc(ws, f"{col}7", f"={col}5-{col}6", border=TOP)
        inp(ws, f"{col}8", 0.05, PCT)
        calc(ws, f"{col}9", f"=ROUND({col}7*{col}8,2)", bold=True)
        calc(ws, f"{col}11", f"={col}9-{col}10")
        calc(ws, f"{col}12", f'=IF({col}11>0,"Expense (increase)",IF({col}11<0,"Income (decrease)","No entry"))', "@")
        calc(ws, f"{col}13", f"=ABS({col}11)", bold=True)
        calc(ws, f"{col}15", f"={col}7")
        calc(ws, f"{col}16", f"={col}9")
        calc(ws, f"{col}17", f"={col}15-{col}16", bold=True, border=TOTAL)
        calc(ws, f"{col}19", f"={col}6")
    inp(ws, "B10", 0)
    calc(ws, "C10", "=B9")
    calc(ws, "D10", "=C9")
    labels = {
        7: "Trade receivables after write-offs",
        8: "Allowance rate",
        9: "Allowance required at year end (step 1)",
        10: "Less: allowance brought forward",
        11: "Change in allowance (step 2)",
        12: "Treatment in Statement of Profit or Loss (step 3)",
        13: "Amount to Statement of Profit or Loss",
        14: "Statement of Financial Position extract",
        15: "Trade receivables",
        16: "Less: allowance for irrecoverable debts",
        17: "Net trade receivables (current asset)",
        18: "Statement of Profit or Loss extract",
        19: "Irrecoverable debts written off (expense)",
    }
    for row, label in labels.items():
        text(ws, f"A{row}", label, bold=row in (9, 14, 17, 18))
    text(ws, "A21", "Rule: only the CHANGE in the allowance goes to the Statement of Profit or Loss.", italic=True)
    text(ws, "A22", "Blue = input. The 2023 opening allowance is 0 because the allowance is created in 2023.", italic=True)

    ws = wb.create_sheet("Allowance Ledger")
    widths(ws, A=10, B=34, C=12, D=3, E=10, F=34, G=12)
    title(ws, "Allowance for irrecoverable debts account",
          "Built from the Allowance Calculator sheet. Amounts of 0 show as a dash.")
    r = 4
    for col, year in (("B", "2023"), ("C", "2024"), ("D", "2025")):
        req = f"'Allowance Calculator'!{col}9"
        opening = f"'Allowance Calculator'!{col}10"
        change = f"'Allowance Calculator'!{col}11"
        text(ws, f"A{r}", f"Year ended 31 December {year}", bold=True)
        header(ws, r + 1, ["Dr", "Details", "$", "", "Cr", "Details", "$"])
        rows = [
            ("Dec 31", "Income statement (decrease)", f"=IF({change}<0,-{change},0)",
             "Jan 1", "Balance b/d", f"={opening}"),
            ("Dec 31", "Balance c/d", f"={req}",
             "Dec 31", "Income statement (increase)", f"=IF({change}>0,{change},0)"),
        ]
        for i, (d1, t1, f1, d2, t2, f2) in enumerate(rows):
            rr = r + 2 + i
            ws[f"A{rr}"], ws[f"B{rr}"] = d1, t1
            ws[f"E{rr}"], ws[f"F{rr}"] = d2, t2
            calc(ws, f"C{rr}", f1)
            calc(ws, f"G{rr}", f2)
        tr = r + 4
        calc(ws, f"C{tr}", f"=C{r+2}+C{r+3}", bold=True, border=TOTAL)
        calc(ws, f"G{tr}", f"=G{r+2}+G{r+3}", bold=True, border=TOTAL)
        ws[f"E{tr+1}"], ws[f"F{tr+1}"] = "Jan 1", "Balance b/d (next year)"
        calc(ws, f"G{tr+1}", f"={req}")
        r = tr + 3

    def info(ws, r):
        header(ws, r, ["Data (from Model Question 8.11)", "2023", "2024", "2025"])
        r += 1
        refs = {}
        rows = [("before", "Trade receivables before write-offs", (40000, 50000, 36000), NUM),
                ("wo", "Debts to be written off", (0, 0, 1000), NUM),
                ("rate", "Allowance rate", (0.03, 0.03, 0.03), PCT)]
        for key, label, vals, fmt in rows:
            text(ws, f"A{r}", label)
            for col, v, y in zip("BCD", vals, ("23", "24", "25")):
                inp(ws, f"{col}{r}", v, fmt)
                refs[f"{key}{y}"] = f"${col}${r}"
            r += 1
        text(ws, f"A{r}", "There was no allowance before 2023.", italic=True)
        return refs, r + 1

    after = {y: f"({{before{y}}}-{{wo{y}}})" for y in ("23", "24", "25")}
    allow = {y: f"ROUND({after[y]}*{{rate{y}}},2)" for y in ("23", "24", "25")}
    change = {"23": allow["23"], "24": f"{allow['24']}-{allow['23']}", "25": f"{allow['25']}-{allow['24']}"}
    lists = '"Expense,Income"'
    qs = ["Part a) Allowance required"]
    for y in ("23", "24", "25"):
        qs.append((f"Allowance required at 31 December 20{y}", allow[y], "num", NUM))
    qs.append("Part b) Effect on the Statement of Profit or Loss")
    for y in ("23", "24", "25"):
        qs.append((f"Amount to SPL in 20{y} (always a positive number)", f"ABS({change[y]})", "num", NUM))
        qs.append((f"Is 20{y} an expense or income?", f'=IF({change[y]}>=0,"Expense","Income")', "text", "@", lists))
    qs.append("Part d) Statement of Financial Position at 31 December 2025")
    qs.append(("Trade receivables (after write-offs)", after["25"], "num", NUM))
    qs.append(("Net trade receivables after deducting the allowance", f"{after['25']}-{allow['25']}", "num", NUM))
    # text templates need refs substituted too, so pre-format them after refs exist
    ws, an = practice(wb, "Practice: Model Question 8.11 (Grace)",
                      ["Grace keeps an allowance for irrecoverable debts. Year end 31 December."],
                      info, qs)
    wb.save(OUT / "Ch08_Irrecoverable_Debts.xlsx")


# ---------------------------------------------------------------- Chapter 10
def build_ch10():
    wb = Workbook()
    read_me(wb, "Chapter 10: Tangible Non-current Assets and Depreciation")

    # Straight line
    ws = wb.create_sheet("Straight Line")
    widths(ws, A=34, B=16, C=16, D=18, E=16)
    title(ws, "Worked example 10.1: straight line method", "Annual depreciation = (Cost - Residual value) / Useful life")
    text(ws, "A4", "Cost of asset ($)")
    inp(ws, "B4", 20000)
    text(ws, "A5", "Residual value ($)")
    inp(ws, "B5", 2000)
    text(ws, "A6", "Useful life (years, up to 10)")
    inp(ws, "B6", 4, "0")
    text(ws, "A7", "Annual depreciation ($)", bold=True)
    calc(ws, "B7", "=(B4-B5)/B6", NUM2, bold=True)
    header(ws, 9, ["Year", "NBV at start", "Depreciation", "Accumulated depreciation", "NBV at end"])
    for i in range(10):
        r = 10 + i
        calc(ws, f"A{r}", f'=IF({i+1}<=$B$6,{i+1},"")', "0")
        calc(ws, f"B{r}", '=IF(A{0}="","",$B$4)'.format(r) if i == 0 else f'=IF(A{r}="","",E{r-1})', NUM2)
        calc(ws, f"C{r}", f'=IF(A{r}="","",$B$7)', NUM2)
        calc(ws, f"D{r}", f'=IF(A{r}="","",C{r})' if i == 0 else f'=IF(A{r}="","",D{r-1}+C{r})', NUM2)
        calc(ws, f"E{r}", f'=IF(A{r}="","",B{r}-C{r})', NUM2)
    text(ws, "A21", "Notice: the depreciation charge is the SAME every year.", italic=True)

    # Reducing balance
    ws = wb.create_sheet("Reducing Balance")
    widths(ws, A=34, B=16, C=16, D=18, E=16)
    title(ws, "Worked example 10.2: reducing balance method", "Depreciation = NBV at start of year x rate")
    text(ws, "A4", "Cost of asset ($)")
    inp(ws, "B4", 10000)
    text(ws, "A5", "Depreciation rate")
    inp(ws, "B5", 0.2, PCT)
    header(ws, 7, ["Year", "NBV at start", "Depreciation", "Accumulated depreciation", "NBV at end"])
    for i in range(6):
        r = 8 + i
        inp(ws, f"A{r}", i + 1, "0")
        ws[f"A{r}"].font = font()
        calc(ws, f"B{r}", "=$B$4" if i == 0 else f"=E{r-1}", NUM2)
        calc(ws, f"C{r}", f"=ROUND(B{r}*$B$5,2)", NUM2)
        calc(ws, f"D{r}", f"=C{r}" if i == 0 else f"=D{r-1}+C{r}", NUM2)
        calc(ws, f"E{r}", f"=B{r}-C{r}", NUM2)
    text(ws, "A15", "Notice: the charge FALLS every year because the NBV falls.", italic=True)
    text(ws, "A16", "Common mistake: using cost instead of NBV after Year 1.", italic=True)

    # Compare
    ws = wb.create_sheet("Compare Methods")
    widths(ws, A=12, B=22, C=22, D=22, E=22)
    title(ws, "Straight line vs reducing balance on the same asset", "Change the blue inputs to compare.")
    text(ws, "A4", "Cost")
    inp(ws, "B4", 16000)
    text(ws, "A5", "Rate")
    inp(ws, "B5", 0.25, PCT)
    text(ws, "C5", "Straight line uses the rate on COST. Reducing balance uses the rate on NBV.", italic=True)
    header(ws, 7, ["Year", "SL depreciation", "SL NBV at end", "RB depreciation", "RB NBV at end"])
    for i in range(4):
        r = 8 + i
        inp(ws, f"A{r}", i + 1, "0")
        ws[f"A{r}"].font = font()
        calc(ws, f"B{r}", "=$B$4*$B$5")
        calc(ws, f"C{r}", f"=$B$4-B{r}" if i == 0 else f"=C{r-1}-B{r}")
        calc(ws, f"D{r}", f"=ROUND($B$4*$B$5,2)" if i == 0 else f"=ROUND(E{r-1}*$B$5,2)", NUM2)
        calc(ws, f"E{r}", f"=$B$4-D{r}" if i == 0 else f"=E{r-1}-D{r}", NUM2)

    # Revaluation
    ws = wb.create_sheet("Revaluation")
    widths(ws, A=44, B=16)
    title(ws, "Worked example 10.3: revaluation method (loose tools)",
          "Depreciation = Opening value + Purchases - Closing value")
    text(ws, "A4", "Loose tools at start of year")
    inp(ws, "B4", 1200)
    text(ws, "A5", "Add: tools bought during the year")
    inp(ws, "B5", 500)
    calc(ws, "B6", "=B4+B5", border=TOP)
    text(ws, "A7", "Less: loose tools valued at end of year")
    inp(ws, "B7", 1100)
    text(ws, "A8", "Depreciation for the year", bold=True)
    calc(ws, "B8", "=B6-B7", bold=True, border=TOTAL)

    # Part year
    ws = wb.create_sheet("Part Year")
    widths(ws, A=48, B=16)
    title(ws, "Section 9: asset bought part way through the year", "Monthly (time) basis")
    text(ws, "A4", "Cost of asset")
    inp(ws, "B4", 12000)
    text(ws, "A5", "Rate per year (straight line on cost)")
    inp(ws, "B5", 0.1, PCT)
    text(ws, "A6", "Months owned in the first year")
    inp(ws, "B6", 9, "0")
    text(ws, "A8", "Full year depreciation")
    calc(ws, "B8", "=B4*B5")
    text(ws, "A9", "Depreciation for months owned (monthly basis)", bold=True)
    calc(ws, "B9", "=B8*B6/12", bold=True)

    # Disposal
    ws = wb.create_sheet("Disposal")
    widths(ws, A=40, B=16, C=3, D=40, E=16)
    title(ws, "Worked example 10.5: disposal of a van", "Profit or loss = Sale proceeds - NBV")
    text(ws, "A4", "Cost of asset")
    inp(ws, "B4", 30000)
    text(ws, "A5", "Accumulated depreciation to date of sale")
    inp(ws, "B5", 18000)
    text(ws, "A6", "Sale proceeds")
    inp(ws, "B6", 10000)
    text(ws, "A8", "NBV at date of sale")
    calc(ws, "B8", "=B4-B5")
    text(ws, "A9", "Profit (+) or loss (-) on disposal", bold=True)
    calc(ws, "B9", "=B6-B8", bold=True)
    calc(ws, "A10", '=IF(B9>0,"Profit on disposal: other income in the SPL",IF(B9<0,"Loss on disposal: expense in the SPL","No profit or loss"))', "@")
    text(ws, "A12", "Disposal account", bold=True)
    header(ws, 13, ["Dr", "$", "", "Cr", "$"])
    ws["A14"] = "Asset account (cost)"
    calc(ws, "B14", "=B4")
    ws["A15"] = "Income statement (profit on disposal)"
    calc(ws, "B15", "=MAX(B9,0)")
    ws["D14"] = "Provision for depreciation"
    calc(ws, "E14", "=B5")
    ws["D15"] = "Bank (sale proceeds)"
    calc(ws, "E15", "=B6")
    ws["D16"] = "Income statement (loss on disposal)"
    calc(ws, "E16", "=MAX(-B9,0)")
    calc(ws, "B17", "=SUM(B14:B16)", bold=True, border=TOTAL)
    calc(ws, "E17", "=SUM(E14:E16)", bold=True, border=TOTAL)

    def info(ws, r):
        return info_rows(ws, r, [
            (None, "Q1 (Model Question 10.2): straight line", None, None),
            ("c1", "Cost", 25000, NUM),
            ("res1", "Residual value", 5000, NUM),
            ("life1", "Useful life (years)", 5, "0"),
            (None, "Q2 (Model Question 10.12): Musa's machine, reducing balance", None, None),
            ("c2", "Cost on 1 January 2023", 40000, NUM),
            ("rate2", "Reducing balance rate", 0.25, PCT),
            ("sale2", "Sale proceeds on 1 January 2026", 15000, NUM),
            (None, "Q3 (Model Question 10.7): loose tools", None, None),
            ("open3", "Value at start of year", 2000, NUM),
            ("buy3", "Bought during year", 800, NUM),
            ("close3", "Value at end of year", 2100, NUM),
            (None, "Q4 (Model Question 10.11): part year", None, None),
            ("c4", "Cost", 12000, NUM),
            ("rate4", "Rate on cost per year", 0.1, PCT),
            ("m4", "Months owned in 2025", 9, "0"),
        ])

    d1 = "{c2}*{rate2}"
    d2 = "({c2}-{c2}*{rate2})*{rate2}"
    d3 = "({c2}-{c2}*{rate2}-({c2}-{c2}*{rate2})*{rate2})*{rate2}"
    nbv3 = f"{{c2}}-({d1})-({d2})-({d3})"
    lists = '"Profit,Loss"'
    qs = [
        "Q1 Straight line",
        ("Annual depreciation", "({c1}-{res1})/{life1}", "num", NUM),
        ("NBV at the end of Year 2", "{c1}-2*({c1}-{res1})/{life1}", "num", NUM),
        "Q2 Reducing balance and disposal",
        ("Depreciation for 2023", d1, "num", NUM),
        ("Depreciation for 2024", d2, "num", NUM),
        ("Depreciation for 2025", d3, "num", NUM),
        ("Accumulated depreciation at 31 December 2025", f"({d1})+({d2})+({d3})", "num", NUM),
        ("NBV at 31 December 2025", nbv3, "num", NUM),
        ("Amount of profit or loss on disposal (positive number)", f"ABS({{sale2}}-({nbv3}))", "num", NUM),
        ("Profit or loss?", f'=IF({{sale2}}-({nbv3})>=0,"Profit","Loss")', "text", "@", lists),
        "Q3 Revaluation",
        ("Depreciation of loose tools", "{open3}+{buy3}-{close3}", "num", NUM),
        "Q4 Part year",
        ("Depreciation for 2025 (monthly basis)", "{c4}*{rate4}*{m4}/12", "num", NUM),
    ]
    practice(wb, "Practice: Depreciation questions",
             ["Use the data below. Show workings on paper, then enter the final figure."], info, qs)
    wb.save(OUT / "Ch10_Depreciation.xlsx")


# ---------------------------------------------------------------- Chapter 11
def build_ch11():
    wb = Workbook()
    read_me(wb, "Chapter 11: Accruals and Prepayments")

    ws = wb.create_sheet("Expense Calculator")
    widths(ws, A=16, B=13, C=14, D=14, E=14, F=14, G=15, H=15, I=15)
    title(ws, "Expense for the year (worked examples 11.1 to 11.3)",
          "Expense = Paid - Opening accrual + Opening prepayment + Closing accrual - Closing prepayment")
    header(ws, 4, ["Expense", "Paid in year", "Opening accrual (owing)", "Opening prepayment",
                   "Closing accrual (owing)", "Closing prepayment", "Expense for SPL",
                   "SFP: other payables", "SFP: other receivables"])
    ws.row_dimensions[4].height = 45
    data = [("Rent", 11000, 0, 0, 1000, 0), ("Insurance", 3600, 0, 0, 0, 900), ("Electricity", 2500, 200, 0, 300, 0)]
    for i, row in enumerate(data):
        r = 5 + i
        ws[f"A{r}"] = row[0]
        for col, v in zip("BCDEF", row[1:]):
            inp(ws, f"{col}{r}", v)
        calc(ws, f"G{r}", f"=B{r}-C{r}+D{r}+E{r}-F{r}", bold=True)
        calc(ws, f"H{r}", f"=E{r}")
        calc(ws, f"I{r}", f"=F{r}")
    text(ws, "A8", "Total", bold=True)
    for col in "BCDEFGHI":
        calc(ws, f"{col}8", f"=SUM({col}5:{col}7)", bold=True, border=TOTAL)
    text(ws, "A10", "Accrued expense: ADD to the expense, show as a current liability.", italic=True)
    text(ws, "A11", "Prepaid expense: SUBTRACT from the expense, show as a current asset.", italic=True)

    ws = wb.create_sheet("Income Calculator")
    widths(ws, A=22, B=13, C=14, D=14, E=14, F=14, G=15, H=15, I=15)
    title(ws, "Income for the year (worked examples 11.4 and 11.5)",
          "Income = Received - Opening accrued + Opening in advance + Closing accrued - Closing in advance")
    header(ws, 4, ["Income", "Received in year", "Opening accrued income (due)", "Opening income in advance",
                   "Closing accrued income (due)", "Closing income in advance", "Income for SPL",
                   "SFP: other receivables", "SFP: other payables"])
    ws.row_dimensions[4].height = 45
    data = [("Rent received", 6500, 0, 0, 0, 500), ("Commission received", 1800, 0, 0, 200, 0)]
    for i, row in enumerate(data):
        r = 5 + i
        ws[f"A{r}"] = row[0]
        for col, v in zip("BCDEF", row[1:]):
            inp(ws, f"{col}{r}", v)
        calc(ws, f"G{r}", f"=B{r}-C{r}+D{r}+E{r}-F{r}", bold=True)
        calc(ws, f"H{r}", f"=E{r}")
        calc(ws, f"I{r}", f"=F{r}")
    text(ws, "A8", "Accrued income: ADD to income, show as a current asset.", italic=True)
    text(ws, "A9", "Income in advance: SUBTRACT from income, show as a current liability.", italic=True)

    ws = wb.create_sheet("Time Apportion")
    widths(ws, A=48, B=16)
    title(ws, "Worked example 11.2: splitting a payment by months", "Insurance paid in advance")
    text(ws, "A4", "Amount paid")
    inp(ws, "B4", 3600)
    text(ws, "A5", "Number of months the payment covers")
    inp(ws, "B5", 12, "0")
    text(ws, "A6", "Months falling in THIS financial year")
    inp(ws, "B6", 9, "0")
    text(ws, "A8", "Cost per month")
    calc(ws, "B8", "=B4/B5", NUM2)
    text(ws, "A9", "Expense for this year (SPL)", bold=True)
    calc(ws, "B9", "=B8*B6", NUM2, bold=True)
    text(ws, "A10", "Prepayment carried to next year (current asset)", bold=True)
    calc(ws, "B10", "=B4-B9", NUM2, bold=True)

    ws = wb.create_sheet("Ledger Account")
    widths(ws, A=10, B=30, C=12, D=3, E=10, F=30, G=12)
    title(ws, "Electricity account (worked example 11.3)",
          "Linked to row 7 of the Expense Calculator. Change the inputs there.")
    header(ws, 4, ["Dr", "Details", "$", "", "Cr", "Details", "$"])
    src = "'Expense Calculator'!"
    ws["A5"], ws["B5"] = "Jan 1", "Balance b/d (prepaid)"
    calc(ws, "C5", f"={src}D7")
    ws["A6"], ws["B6"] = "", "Bank (paid during year)"
    calc(ws, "C6", f"={src}B7")
    ws["A7"], ws["B7"] = "Dec 31", "Balance c/d (owing)"
    calc(ws, "C7", f"={src}E7")
    ws["E5"], ws["F5"] = "Jan 1", "Balance b/d (owing)"
    calc(ws, "G5", f"={src}C7")
    ws["E6"], ws["F6"] = "Dec 31", "Income statement"
    calc(ws, "G6", f"={src}G7")
    ws["E7"], ws["F7"] = "Dec 31", "Balance c/d (prepaid)"
    calc(ws, "G7", f"={src}F7")
    calc(ws, "C8", "=SUM(C5:C7)", bold=True, border=TOTAL)
    calc(ws, "G8", "=SUM(G5:G7)", bold=True, border=TOTAL)
    ws["E9"], ws["F9"] = "Jan 1", "Balance b/d (owing)"
    calc(ws, "G9", f"={src}E7")
    ws["A9"], ws["B9"] = "Jan 1", "Balance b/d (prepaid)"
    calc(ws, "C9", f"={src}F7")
    text(ws, "A11", "Both totals must agree. The Income statement figure is the balancing figure.", italic=True)

    def info(ws, r):
        return info_rows(ws, r, [
            (None, "Model Question 11.11 (Zara, year ended 31 March 2026)", None, None),
            ("wo", "Wages owing at 1 April 2025", 300, NUM),
            ("wp", "Wages paid during the year", 18000, NUM),
            ("wc", "Wages owing at 31 March 2026", 450, NUM),
            ("ro", "Rates prepaid at 1 April 2025", 200, NUM),
            ("rp", "Rates paid during the year", 1600, NUM),
            ("rc", "Rates prepaid at 31 March 2026", 250, NUM),
            ("cr", "Commission received during the year", 3000, NUM),
            ("cd", "Commission still due at 31 March 2026", 250, NUM),
            (None, "Model Question 11.7", None, None),
            ("ip", "Insurance paid on 1 October 2025 for 12 months", 4800, NUM),
            ("im", "Months in the year to 31 December 2025", 3, "0"),
        ])

    qs = [
        "Model Question 11.11",
        ("Wages expense for the SPL", "{wp}-{wo}+{wc}", "num", NUM),
        ("Rates expense for the SPL", "{rp}+{ro}-{rc}", "num", NUM),
        ("Commission receivable for the SPL", "{cr}+{cd}", "num", NUM),
        ("Wages account: total of each side", "{wp}+{wc}", "num", NUM),
        ("SFP: total other receivables", "{rc}+{cd}", "num", NUM),
        ("SFP: total other payables", "{wc}", "num", NUM),
        "Model Question 11.7",
        ("Insurance expense for 2025", "{ip}/12*{im}", "num", NUM),
        ("Insurance prepaid at 31 December 2025", "{ip}-{ip}/12*{im}", "num", NUM),
    ]
    practice(wb, "Practice: Accruals and prepayments", ["Use the data below."], info, qs)
    wb.save(OUT / "Ch11_Accruals_Prepayments.xlsx")


# ---------------------------------------------------------------- Chapter 12
CONCEPTS = [
    ("Business entity", "The business is separate from its owner.", "Owner's private house is not recorded; goods taken are drawings."),
    ("Duality", "Every transaction has two equal effects: a debit and a credit.", "Buy equipment by cheque: equipment up, bank down."),
    ("Money measurement", "Only record items with a money value.", "Staff skill and customer loyalty are not recorded."),
    ("Historical cost", "Record assets at their original cost.", "Land bought for $50,000 stays at $50,000."),
    ("Going concern", "Assume the business will continue trading.", "Assets at cost less depreciation, not closing-down value."),
    ("Realisation", "Record revenue when it is earned (goods pass to customer).", "Credit sale on 28 Dec is this year's revenue."),
    ("Accruals", "Match income and expenses to the period they belong to.", "Accrued electricity; depreciation."),
    ("Consistency", "Use the same methods every year.", "Keep the same depreciation method."),
    ("Prudence", "Do not overstate assets or profit.", "Allowance for irrecoverable debts; inventory at lower of cost and NRV."),
    ("Materiality", "Small items can be treated simply.", "A $5 stapler is an expense, not a non-current asset."),
]


def build_ch12():
    wb = Workbook()
    read_me(wb, "Chapter 12: Fundamental Accounting Principles and Concepts",
            ["The Quiz sheet uses drop-down lists. Click a yellow cell and pick a concept."])

    ws = wb.create_sheet("Concepts")
    widths(ws, A=22, B=55, C=65)
    title(ws, "The ten accounting concepts")
    header(ws, 3, ["Concept", "Meaning", "Example"])
    for i, (name, meaning, example) in enumerate(CONCEPTS):
        r = 4 + i
        text(ws, f"A{r}", name, bold=True)
        text(ws, f"B{r}", meaning, wrap=True)
        text(ws, f"C{r}", example, wrap=True)
    r = 4 + len(CONCEPTS) + 1
    text(ws, f"A{r}", "Qualities of useful information", bold=True)
    for j, (q, m) in enumerate([
        ("Relevance", "Helps users make decisions; available in time."),
        ("Reliability", "Free from errors and bias; can be checked."),
        ("Comparability", "Can be compared with other years and businesses."),
        ("Understandability", "Clear to users with reasonable knowledge."),
    ]):
        text(ws, f"A{r+1+j}", q, bold=True)
        text(ws, f"B{r+1+j}", m)

    ws = wb.create_sheet("NRV Calculator")
    widths(ws, A=44, B=16)
    title(ws, "Inventory: lower of cost and net realisable value (prudence)")
    text(ws, "A4", "Cost of goods")
    inp(ws, "B4", 800)
    text(ws, "A5", "Expected selling price")
    inp(ws, "B5", 700)
    text(ws, "A6", "Costs to sell (repairs, selling costs)")
    inp(ws, "B6", 50)
    text(ws, "A8", "Net realisable value (NRV)")
    calc(ws, "B8", "=B5-B6")
    text(ws, "A9", "Value for inventory (lower of cost and NRV)", bold=True)
    calc(ws, "B9", "=MIN(B4,B8)", bold=True, border=TOTAL)

    lst = wb.create_sheet("Lists")
    for i, (name, _, _) in enumerate(CONCEPTS):
        lst[f"A{i+1}"] = name
    lst.sheet_state = "hidden"
    options = f"Lists!$A$1:$A${len(CONCEPTS)}"

    scenarios = [
        ("12.1 The owner's private car is not in the business accounts.", "Business entity"),
        ("12.2 Land is recorded at the price paid, not its current value.", "Historical cost"),
        ("12.3 The same depreciation method is used every year.", "Consistency"),
        ("12.4 A $5 calculator is recorded as an expense.", "Materiality"),
        ("12.5 Revenue is recorded when goods are delivered.", "Realisation"),
        ("12.6 Staff skill is not recorded in the accounts.", "Money measurement"),
        ("12.7 Every transaction has a debit and a credit.", "Duality"),
        ("12.8 Assets shown at cost less depreciation, not closing-down prices.", "Going concern"),
        ("12.9 December electricity paid in January is December's expense.", "Accruals"),
        ("12.10 An allowance for irrecoverable debts is created.", "Prudence"),
        ("Extra: Inventory is valued at the lower of cost and NRV.", "Prudence"),
        ("Extra: Goods taken by the owner are recorded as drawings.", "Business entity"),
        ("Extra: A credit sale on 30 Dec, paid 10 Jan, is this year's revenue.", "Realisation"),
        ("Extra: Insurance prepaid is deducted from this year's expense.", "Accruals"),
        ("Extra: Equipment $5,000 by cheque: equipment up, bank down.", "Duality"),
    ]
    qs = ["Pick the concept for each situation"] + [(s, f'="{a}"', "text", "@", options) for s, a in scenarios]
    ws, _ = practice(wb, "Quiz: Which concept applies?", ["Choose from the drop-down list in each yellow cell."],
                     lambda ws, r: ({}, r), qs)
    ws.title = "Quiz"
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 22
    wb.save(OUT / "Ch12_Accounting_Concepts.xlsx")


# ---------------------------------------------------------------- Financial statements
AMINA_TB = [
    ("rev", "Revenue", None, 110000),
    ("pur", "Purchases", 62000, None),
    ("sr", "Sales returns", 1500, None),
    ("pr", "Purchases returns", None, 1200),
    ("ci", "Carriage inwards", 800, None),
    ("co", "Carriage outwards", 600, None),
    ("oi", "Inventory at 1 January 2025", 8000, None),
    ("wag", "Wages", 14000, None),
    ("rent", "Rent", 5500, None),
    ("ins", "Insurance", 1800, None),
    ("gen", "General expenses", 2100, None),
    ("da", "Discount allowed", 400, None),
    ("dr", "Discount received", None, 500),
    ("id", "Irrecoverable debts", 700, None),
    ("prem", "Premises at cost", 60000, None),
    ("eq", "Equipment at cost", 20000, None),
    ("eqd", "Provision for depreciation: equipment", None, 6000),
    ("tr", "Trade receivables", 12000, None),
    ("afd", "Allowance for irrecoverable debts", None, 400),
    ("tp", "Trade payables", None, 7500),
    ("bank", "Bank", 4300, None),
    ("cash", "Cash", 200, None),
    ("loan", "Loan (repayable 2030)", None, 10000),
    ("cap", "Capital", None, 67300),
    ("draw", "Drawings", 9000, None),
]


def write_tb(ws, r, rows):
    """Write a trial balance. Returns dict key -> absolute address of its amount."""
    header(ws, r, ["Account", "Dr $", "Cr $"])
    r += 1
    first = r
    refs = {}
    for key, name, dr, cr in rows:
        ws[f"A{r}"] = name
        ws[f"A{r}"].font = font()
        if dr is not None:
            inp(ws, f"B{r}", dr)
            refs[key] = f"$B${r}"
        else:
            inp(ws, f"C{r}", cr)
            refs[key] = f"$C${r}"
        r += 1
    text(ws, f"A{r}", "Totals", bold=True)
    calc(ws, f"B{r}", f"=SUM(B{first}:B{r-1})", bold=True, border=TOTAL)
    calc(ws, f"C{r}", f"=SUM(C{first}:C{r-1})", bold=True, border=TOTAL)
    calc(ws, f"D{r}", f'=IF(B{r}=C{r},"Trial balance agrees","Does not agree")', "@")
    return refs, r + 1


def build_fs():
    wb = Workbook()
    read_me(wb, "Statement of Profit or Loss and Statement of Financial Position",
            ["Worked example: Amina Stores (notes section 6). Practice: Ben's Bikes (Model Question FS.9)."])

    tb = wb.create_sheet("Trial Balance")
    widths(tb, A=42, B=14, C=14, D=24)
    title(tb, "Amina Stores: Trial balance at 31 December 2025")
    T, _ = write_tb(tb, 3, AMINA_TB)
    T = {k: f"'Trial Balance'!{v}" for k, v in T.items()}

    adj = wb.create_sheet("Adjustments")
    widths(adj, A=62, B=14)
    title(adj, "Amina Stores: Additional information at 31 December 2025")
    A, _ = info_rows(adj, 3, [
        ("ci", "1. Closing inventory", 9500, NUM),
        ("ra", "2. Rent owing (accrued)", 500, NUM),
        ("ip", "3. Insurance prepaid", 300, NUM),
        ("ar", "4. Allowance for irrecoverable debts (% of trade receivables)", 0.05, PCT),
        ("dr", "5. Equipment depreciation (% of cost, straight line)", 0.1, PCT),
        ("lr", "6. Loan interest rate per year (none paid)", 0.05, PCT),
    ])
    A = {k: f"Adjustments!{v}" for k, v in A.items()}

    spl = wb.create_sheet("SPL")
    widths(spl, A=54, B=14, C=14)
    title(spl, "Amina Stores: Statement of Profit or Loss for the year ended 31 December 2025")
    header(spl, 3, ["", "$", "$"])
    R = {}
    row = [4]

    def line(key, label, col, formula, bold=False, border=None):
        r = row[0]
        text(spl, f"A{r}", label, bold=bold)
        if formula:
            calc(spl, f"{col}{r}", formula, bold=bold, border=border)
        R[key] = f"{col}{r}"
        row[0] += 1

    line("rev", "Revenue", "C", f"={T['rev']}")
    line("sr", "Less: Sales returns", "C", f"={T['sr']}")
    line("nr", "Net revenue", "C", f"={R['rev']}-{R['sr']}", True, TOP)
    line("h1", "Less: Cost of sales", "B", None, True)
    line("oi", "Opening inventory", "B", f"={T['oi']}")
    line("pur", "Add: Purchases", "B", f"={T['pur']}")
    line("pr", "Less: Purchases returns", "B", f"=-{T['pr']}")
    line("ci", "Add: Carriage inwards", "B", f"={T['ci']}")
    line("avail", "Cost of goods available for sale", "B", f"=SUM({R['oi']}:{R['ci']})", False, TOP)
    line("cl", "Less: Closing inventory", "B", f"=-{A['ci']}")
    line("cos", "Cost of sales", "C", f"={R['avail']}+{R['cl']}", True)
    line("gp", "Gross profit", "C", f"={R['nr']}-{R['cos']}", True, TOP)
    line("h2", "Add: Other income", "B", None, True)
    line("dr", "Discount received", "B", f"={T['dr']}")
    allow_change = f"ROUND({T['tr']}*{A['ar']},2)-{T['afd']}"
    line("dec", "Decrease in allowance for irrecoverable debts", "B", f"=MAX(-({allow_change}),0)")
    line("toi", "Total other income", "C", f"={R['dr']}+{R['dec']}")
    line("sub", "", "C", f"={R['gp']}+{R['toi']}", False, TOP)
    line("h3", "Less: Expenses", "B", None, True)
    first_exp = row[0]
    line("wag", "Wages", "B", f"={T['wag']}")
    line("rent", "Rent (paid + owing)", "B", f"={T['rent']}+{A['ra']}")
    line("ins", "Insurance (paid - prepaid)", "B", f"={T['ins']}-{A['ip']}")
    line("gen", "General expenses", "B", f"={T['gen']}")
    line("co", "Carriage outwards", "B", f"={T['co']}")
    line("da", "Discount allowed", "B", f"={T['da']}")
    line("id", "Irrecoverable debts", "B", f"={T['id']}")
    line("inc", "Increase in allowance for irrecoverable debts", "B", f"=MAX({allow_change},0)")
    line("dep", "Depreciation: equipment", "B", f"={T['eq']}*{A['dr']}")
    line("li", "Loan interest", "B", f"={T['loan']}*{A['lr']}")
    last_exp = row[0] - 1
    line("texp", "Total expenses", "C", f"=SUM(B{first_exp}:B{last_exp})", False, TOP)
    line("np", "Profit for the year", "C", f"={R['sub']}-{R['texp']}", True, TOTAL)
    SPL = {k: f"SPL!{v}" for k, v in R.items()}

    sfp = wb.create_sheet("SFP")
    widths(sfp, A=44, B=14, C=16, D=14)
    title(sfp, "Amina Stores: Statement of Financial Position at 31 December 2025")
    header(sfp, 3, ["", "Cost $", "Acc. dep. $", "NBV $"])
    S = {}
    srow = [4]

    def sline(key, label, cells, bold=False, border=None):
        r = srow[0]
        text(sfp, f"A{r}", label, bold=bold)
        for col, formula in cells.items():
            calc(sfp, f"{col}{r}", formula, bold=bold, border=border)
            S[key + col] = f"{col}{r}"
        srow[0] += 1

    sline("h", "Non-current assets", {}, True)
    sline("prem", "Premises", {"B": f"={T['prem']}", "C": "=0", "D": "=B{0}-C{0}".format(srow[0])})
    sline("eq", "Equipment", {"B": f"={T['eq']}", "C": f"={T['eqd']}+{SPL['dep']}", "D": "=B{0}-C{0}".format(srow[0])})
    sline("nca", "Total non-current assets", {
        "B": f"={S['premB']}+{S['eqB']}", "C": f"={S['premC']}+{S['eqC']}", "D": f"={S['premD']}+{S['eqD']}"}, True, TOP)
    sline("h2", "Current assets", {}, True)
    sline("inv", "Inventory", {"C": f"={A['ci']}"})
    sline("tr", "Trade receivables", {"B": f"={T['tr']}"})
    sline("al", "Less: Allowance for irrecoverable debts", {"B": f"=ROUND({T['tr']}*{A['ar']},2)", "C": "=B{0}-B{1}".format(srow[0] - 1, srow[0])})
    sline("or", "Other receivables (insurance prepaid)", {"C": f"={A['ip']}"})
    sline("bank", "Bank", {"C": f"={T['bank']}"})
    sline("cash", "Cash", {"C": f"={T['cash']}"})
    sline("tca", "Total current assets", {"C": f"=SUM({S['invC']}:{S['cashC']})"}, True, TOP)
    sline("h3", "Current liabilities", {}, True)
    sline("tp", "Trade payables", {"C": f"={T['tp']}"})
    sline("op", "Other payables (rent owing + loan interest owing)", {"C": f"={A['ra']}+{SPL['li']}"})
    sline("tcl", "Total current liabilities", {"C": f"={S['tpC']}+{S['opC']}"}, True, TOP)
    sline("ncur", "Net current assets (working capital)", {"D": f"={S['tcaC']}-{S['tclC']}"}, True)
    sline("sub", "", {"D": f"={S['ncaD']}+{S['ncurD']}"}, False, TOP)
    sline("h4", "Non-current liabilities", {}, True)
    sline("loan", "Loan", {"D": f"={T['loan']}"})
    sline("na", "Net assets", {"D": f"={S['subD']}-{S['loanD']}"}, True, TOTAL)
    srow[0] += 1
    sline("h5", "Capital", {}, True)
    sline("oc", "Opening capital", {"D": f"={T['cap']}"})
    sline("np", "Add: Profit for the year", {"D": f"={SPL['np']}"})
    sline("s2", "", {"D": f"={S['ocD']}+{S['npD']}"}, False, TOP)
    sline("dw", "Less: Drawings", {"D": f"={T['draw']}"})
    sline("cc", "Closing capital", {"D": f"={S['s2D']}-{S['dwD']}"}, True, TOTAL)
    srow[0] += 1
    sline("chk", "Check", {"D": f'=IF(ROUND({S["naD"]}-{S["ccD"]},2)=0,"Balances","Does not balance")'}, True)

    # Practice: Ben's Bikes
    ben_tb = [
        ("rev", "Revenue", None, 82000),
        ("pur", "Purchases", 45000, None),
        ("sr", "Sales returns", 800, None),
        ("pr", "Purchases returns", None, 900),
        ("ci", "Carriage inwards", 400, None),
        ("oi", "Inventory at 1 July 2025", 6200, None),
        ("wag", "Wages", 12600, None),
        ("elec", "Electricity", 2300, None),
        ("ins", "Insurance", 1600, None),
        ("adv", "Advertising", 1200, None),
        ("id", "Irrecoverable debts", 350, None),
        ("rr", "Rent received", None, 2400),
        ("mv", "Motor vehicles at cost", 24000, None),
        ("mvd", "Provision for depreciation: motor vehicles", None, 6000),
        ("fx", "Fixtures at cost", 8000, None),
        ("fxd", "Provision for depreciation: fixtures", None, 2000),
        ("tr", "Trade receivables", 9000, None),
        ("afd", "Allowance for irrecoverable debts", None, 500),
        ("tp", "Trade payables", None, 5300),
        ("bank", "Bank", 3150, None),
        ("cap", "Capital", None, 23000),
        ("draw", "Drawings", 7500, None),
    ]

    def info(ws, r):
        text(ws, f"A{r}", "Trial balance at 30 June 2026", bold=True)
        refs, r = write_tb(ws, r + 1, ben_tb)
        r += 1
        text(ws, f"A{r}", "Additional information at 30 June 2026", bold=True)
        more, r = info_rows(ws, r + 1, [
            ("cinv", "1. Closing inventory", 7100, NUM),
            ("eo", "2. Electricity owing", 250, NUM),
            ("ip", "3. Insurance prepaid", 400, NUM),
            ("ria", "4. Rent received in advance", 200, NUM),
            ("ar", "5. Allowance (% of trade receivables)", 0.04, PCT),
            ("mvr", "6. Motor vehicles depreciation (reducing balance)", 0.2, PCT),
            ("fxr", "7. Fixtures depreciation (straight line on cost)", 0.1, PCT),
        ])
        refs.update(more)
        return refs, r

    nr = "({rev}-{sr})"
    cos = "({oi}+{pur}-{pr}+{ci}-{cinv})"
    gp = f"({nr}-{cos})"
    newal = "ROUND({tr}*{ar},2)"
    ch = f"({newal}-{{afd}})"
    rent = "({rr}-{ria})"
    dec = f"MAX(-{ch},0)"
    inc = f"MAX({ch},0)"
    dmv = "(({mv}-{mvd})*{mvr})"
    dfx = "({fx}*{fxr})"
    texp = f"({{wag}}+{{elec}}+{{eo}}+{{ins}}-{{ip}}+{{adv}}+{{id}}+{inc}+{dmv}+{dfx})"
    profit = f"({gp}+{rent}+{dec}-{texp})"
    nca = f"({{mv}}-{{mvd}}-{dmv}+{{fx}}-{{fxd}}-{dfx})"
    tca = f"({{cinv}}+{{tr}}-{newal}+{{ip}}+{{bank}})"
    tcl = "({tp}+{eo}+{ria})"
    qs = [
        "a) Statement of Profit or Loss",
        ("Net revenue", nr, "num", NUM),
        ("Cost of goods available for sale (before closing inventory)", "{oi}+{pur}-{pr}+{ci}", "num", NUM),
        ("Cost of sales", cos, "num", NUM),
        ("Gross profit", gp, "num", NUM),
        ("Rent received (income for the year)", rent, "num", NUM),
        ("Decrease in allowance for irrecoverable debts", dec, "num", NUM),
        ("Electricity expense", "{elec}+{eo}", "num", NUM),
        ("Insurance expense", "{ins}-{ip}", "num", NUM),
        ("Depreciation: motor vehicles", dmv, "num", NUM),
        ("Depreciation: fixtures", dfx, "num", NUM),
        ("Total expenses", texp, "num", NUM),
        ("Profit for the year", profit, "num", NUM),
        "b) Statement of Financial Position",
        ("Motor vehicles: NBV", f"{{mv}}-{{mvd}}-{dmv}", "num", NUM),
        ("Fixtures: NBV", f"{{fx}}-{{fxd}}-{dfx}", "num", NUM),
        ("Total non-current assets", nca, "num", NUM),
        ("Net trade receivables (after allowance)", f"{{tr}}-{newal}", "num", NUM),
        ("Total current assets", tca, "num", NUM),
        ("Other payables", "{eo}+{ria}", "num", NUM),
        ("Total current liabilities", tcl, "num", NUM),
        ("Net current assets (working capital)", f"{tca}-{tcl}", "num", NUM),
        ("Net assets", f"{nca}+{tca}-{tcl}", "num", NUM),
        ("Closing capital", f"{{cap}}+{profit}-{{draw}}", "num", NUM),
    ]
    ws, _ = practice(wb, "Practice: Ben's Bikes (Model Question FS.9)",
                     ["Prepare both statements on paper, then enter the key figures below."], info, qs)
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 24
    wb.save(OUT / "Financial_Statements.xlsx")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    build_ch08()
    build_ch10()
    build_ch11()
    build_ch12()
    build_fs()
    print("Workbooks written to", OUT)
