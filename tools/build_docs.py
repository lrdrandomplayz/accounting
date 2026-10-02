"""Fill the generated sections of the Markdown files from the workbooks.

Run after tools/build_excel.py and after the workbooks have been recalculated
(open and save in Excel, or LibreOffice). If the workbooks have no cached values,
this script recalculates temporary copies with LibreOffice (soffice).

Markers in the Markdown files:
  <!-- BEGIN GENERATED: questions ch08 --> ... <!-- END GENERATED: questions ch08 -->
  <!-- BEGIN GENERATED: answers ch08 -->   ... <!-- END GENERATED: answers ch08 -->
  <!-- BEGIN GENERATED: ch08.ex1.journal --> ... (a tagged worked-example block)
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
from itertools import zip_longest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
LEVELS = {1: "Level 1: Easy", 2: "Level 2: Easy to medium", 3: "Level 3: Medium", 4: "Level 4: Hard",
          5: "Level 5: Challenge"}
PCT, INT, NUM2 = "0.0%", "0", '#,##0.00;(#,##0.00);"-"'
YEAR = "@Y"


MD = {
    "en": {"item": "Item", "figure": "Figure", "answer": "Answer", "part": "Particulars", "dr_rm": "Debit (RM)",
           "cr_rm": "Credit (RM)", "total": "Total", "date": "Date", "debit": "Debit", "credit": "Credit",
           "dr": "Dr", "cr": "Cr", "folio": "Folio", "amount": "Amount", "cost": "Cost",
           "accdep": "Accumulated Depreciation", "carrying": "Carrying Amount", "question": "Question",
           "level": "Level"},
    "ms": {"item": "Item", "figure": "Angka", "answer": "Jawapan", "part": "Butir", "dr_rm": "Debit (RM)",
           "cr_rm": "Kredit (RM)", "total": "Jumlah", "date": "Tarikh", "debit": "Debit", "credit": "Kredit",
           "dr": "Dt", "cr": "Kt", "folio": "Folio", "amount": "Amaun", "cost": "Kos",
           "accdep": "Susut Nilai Terkumpul", "carrying": "Nilai Buku", "question": "Soalan", "level": "Tahap"},
}
L = MD["en"]


def load_values(fname, folder="excel"):
    path = ROOT / folder / fname
    wb = load_workbook(path, data_only=True)
    formulas = load_workbook(path)
    cells = [(ws.title, c.coordinate) for ws in formulas.worksheets for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and c.value.startswith("=")]
    empty = sum(1 for t, a in cells if wb[t][a].value is None)
    if empty < len(cells) / 2:   # blank-by-design cells (e.g. unanswered checks) read as None
        return wb
    if False:
        return wb
    print(f"  {fname}: no cached values, recalculating a copy with LibreOffice")
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "in" / fname
        src.parent.mkdir()
        shutil.copy(path, src)
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(src)],
                       check=True, capture_output=True)
        return load_workbook(Path(tmp) / fname, data_only=True)


def money(v, fmt=None, zero="-"):
    if v is None or v == "":
        return ""
    if isinstance(v, str):
        return v
    if fmt == PCT:
        return f"{v * 100:g}%"
    if fmt == INT:
        return f"{v:g}"
    if abs(v) < 1e-9:
        return zero
    if fmt == NUM2 or abs(v - round(v)) > 1e-6:
        s = f"{abs(v):,.2f}"
    else:
        s = f"{abs(round(v)):,}"
    return f"({s})" if v < 0 else s


def cell(get, a, fmt=None):
    return money(get(a), fmt)


def esc(t):
    return str(t).replace("|", "\\|")


# ------------------------------------------------------------------ block renderers
def r_data(b):
    out = [f"| {L['item']} | {L['figure']} |", "|---|---|"]
    for label, value, fmt in b["rows"]:
        if label == "#":
            out.append(f"| **{esc(value)}** | |")
        else:
            out.append(f"| {esc(label)} | {money(value, fmt, zero='0')} |")
    return "\n".join(out)


def r_tb(b):
    out = [f"**{b['title']}**", "", f"| {L['part']} | {L['dr_rm']} | {L['cr_rm']} |", "|---|---:|---:|"]
    for label, dr, cr in b["rows"]:
        out.append(f"| {esc(label)} | {money(dr) if dr is not None else ''} | {money(cr) if cr is not None else ''} |")
    d, c = b["totals"]
    out.append(f"| **{L['total']}** | **{money(d)}** | **{money(c)}** |")
    return "\n".join(out)


def r_calc(b, get):
    out = [f"| {L['item']} | {L['answer']} |", "|---|---:|"]
    for label, a, fmt in b["rows"]:
        out.append(f"| {esc(label)} | **{cell(get, a, fmt)}** |")
    return "\n".join(out)


def r_choice(b, get):
    out = [f"| {L['item']} | {L['answer']} |", "|---|---|"]
    for label, a, _acc in b["rows"]:
        out.append(f"| {esc(label)} | **{esc(get(a))}** |")
    return "\n".join(out)


def r_written(b):
    return "\n\n".join(line.strip() for line in b["text"].split("\n") if line.strip())


def r_table(b, get):
    heads = b["headers"]
    out = ["| " + " | ".join(esc(h) for h in heads) + " |", "|---" + "|---:" * (len(heads) - 1) + "|"]
    for label, cells in b["rows"]:
        vals = []
        values = []
        for c in cells:
            if c is None:
                vals.append("")
            elif c[0] == "v":
                v = get(c[1])
                values.append(v)
                vals.append(money(v, c[2]))
            elif c[0] == "d":
                vals.append(money(c[1], c[2], zero="0"))
            else:
                vals.append(esc(c[1]))
        if values and all(v in (None, "") for v in values):
            continue
        out.append(f"| {esc(label)} | " + " | ".join(vals) + " |")
    return "\n".join(out)


def r_journal(b, get):
    out = [f"| {L['date']} | {L['part']} | {L['debit']} | {L['credit']} |", "|---|---|---:|---:|"]
    for row in b["rows"]:
        if row[0] == "y":
            rm = "**RM**" if row[2] else ""
            out.append(f"| **{row[1]}** | | {rm} | {rm} |")
        elif row[0] == "n":
            out.append(f"| | *({esc(row[1])})* | | |")
        elif row[0] == "d":
            out.append(f"| {row[1]} | {esc(row[2])} | {cell(get, row[3])} | |")
        else:
            out.append(f"| {row[1]} | &emsp;&emsp;{esc(row[2])} | | {cell(get, row[3])} |")
    return "\n".join(out)


def r_ledger(b, get):
    """Ledger in Date | Particulars | Folio | Amount format, with the year (and RM) on its own row.

    Lines whose amount is nil in a worked example are left out, and year rows left empty are dropped.
    """
    side = f"{L['date']} | {L['part']} | {L['folio']} | {L['amount']}"
    out = [f"**{L['dr']}** &emsp;&emsp; **{b['name']}** &emsp;&emsp; **{L['cr']}**", "",
           f"| {side} | {side} |",
           "|---|---|---|---:|---|---|---|---:|"]
    sides = {"dr": [], "cr": []}
    rm_done = {"dr": False, "cr": False}

    def show(e, side):
        if e is None:
            return " | | | "
        if e[0] == YEAR:
            rm = ""
            if not rm_done[side]:
                rm, rm_done[side] = "**RM**", True
            return f"**{e[1]}** | | | {rm}"
        return f"{e[0]} | {esc(e[1])} | {e[2]} | {money(get(e[3]))}"

    def tidy(items):
        kept = []
        for i, e in enumerate(items):
            if e[0] == YEAR and (i + 1 == len(items) or items[i + 1][0] == YEAR):
                continue
            kept.append(e)
        return kept

    def flush(total=None):
        for d, c in zip_longest(tidy(sides["dr"]), tidy(sides["cr"])):
            left, right = show(d, "dr"), show(c, "cr")
            out.append(f"| {left} | {right} |")
        if total:
            out.append(f"| | | | **{money(get(total['dr']))}** | | | | **{money(get(total['cr']))}** |")
        sides["dr"].clear()
        sides["cr"].clear()

    for row in b["rows"]:
        if row.get("total"):
            flush(row)
            continue
        for side in ("dr", "cr"):
            e = row.get(side)
            if not e:
                continue
            if e[0] == YEAR or get(e[3]) not in (None, "", 0, 0.0):
                sides[side].append(e)
    flush()
    return "\n".join(out)


def r_statement(b, get):
    out = ["  \n".join(f"**{h}**" for h in b["heading"]), "",
           f"| | {L['part']} | RM | RM | RM |", "|---|---|---:|---:|---:|"]
    if b["kind"] == "sofp":
        out.append(f"| | | **{L['cost']}** | **{L['accdep']}** | **{L['carrying']}** |")
    for m, p, vals, flags in b["rows"]:
        name = f"**{esc(p)}**" if p and ("h" in flags or "b" in flags) else esc(p)
        cols = []
        for c in ("1", "2", "3"):
            if c in vals:
                a, line = vals[c]
                s = cell(get, a)
                cols.append(f"**{s}**" if line == "total" or "b" in flags else s)
            else:
                cols.append("")
        out.append(f"| {m} | {name} | " + " | ".join(cols) + " |")
    return "\n".join(out)


def r_mcq_question(b):
    out = []
    for n, (stem, options, _a) in enumerate(b["rows"], 1):
        opts = "\n".join(f"* {x}. {esc(o)}" for x, o in zip("ABCD", options))
        out.append(f"**{n}.** {stem}\n\n{opts}")
    return "\n\n".join(out)


def r_mcq(b, get):
    out = [f"| {L['question']} | {L['answer']} |", "|---|---|"]
    for n, (_stem, options, a) in enumerate(b["rows"], 1):
        letter = get(a)
        text = options["ABCD".index(letter)] if letter in ("A", "B", "C", "D") else ""
        out.append(f"| {n} | **{letter}** ({esc(text)}) |")
    return "\n".join(out)


def render_answer_block(b, get):
    k = b["k"]
    if k == "mcq":
        return r_mcq(b, get)
    if k == "part":
        return f"**{b['text']}**"
    if k == "calc":
        return r_calc(b, get)
    if k == "choice":
        return r_choice(b, get)
    if k == "written":
        return r_written(b)
    if k == "table":
        return r_table(b, get) if any(c and c[0] == "v" for _, cells in b["rows"] for c in cells) else None
    if k == "journal":
        return r_journal(b, get)
    if k == "ledger":
        return r_ledger(b, get)
    if k == "statement":
        return r_statement(b, get)
    return None


def render_question(sheet):
    parts = [f"**{sheet['label']}**" if sheet["section"] == "pa" else f"**{sheet['qid']} ({sheet['label']})**"]
    bullets = []

    def flush_bullets():
        if bullets:
            parts.append("\n".join(f"* {x}" for x in bullets))
            bullets.clear()

    for b in sheet["blocks"]:
        k = b["k"]
        if k == "text" and b["style"] == "bullet":
            bullets.append(b["text"])
            continue
        flush_bullets()
        if k == "text":
            parts.append(b["text"])
        elif k == "data" and b["md"]:
            parts.append(r_data(b))
        elif k == "tb":
            parts.append(r_tb(b))
        elif k == "table" and not any(c and c[0] == "v" for _, cells in b["rows"] for c in cells):
            parts.append(r_table(b, lambda a: None))
        elif k == "table" and sheet["section"] == "pa" and all(re.match(r"^\d+\. ", lab) for lab, _c in b["rows"]):
            parts.append("\n".join(esc(label) for label, _cells in b["rows"]))
        elif k == "part":
            parts.append(b["text"])
        elif k == "mcq":
            parts.append(r_mcq_question(b))
        elif k == "choice":
            subs = [label for label, _a, _acc in b["rows"] if re.match(r"^[a-z]\) ", label)]
            if subs:
                parts.append("\n\n".join(subs))
    flush_bullets()
    return "\n\n".join(parts)


def render_answers(sheet, get):
    head = f"Level {sheet['level']}" if sheet["section"] == "structured" else sheet["label"]
    parts = [f"### {sheet['label']}" if sheet["section"] == "pa" else f"### {sheet['qid']} ({head})"]
    for b in sheet["blocks"]:
        s = render_answer_block(b, get)
        if s:
            parts.append(s)
    return "\n\n".join(parts)


def replace_section(text, name, body, path):
    pat = re.compile(rf"(<!-- BEGIN GENERATED: {re.escape(name)} -->).*?(<!-- END GENERATED: {re.escape(name)} -->)",
                     re.S)
    if not pat.search(text):
        return text, False
    return pat.sub(lambda m: m.group(1) + "\n" + body + "\n" + m.group(2), text), True


def main():
    global L
    manifest = json.loads((ROOT / "tools" / "manifest.json").read_text())
    pa = ROOT / "tools" / "manifest_pa.json"
    if pa.exists():
        manifest.update(json.loads(pa.read_text()))
    sections = {}
    for fname, book in manifest.items():
        L = MD[book.get("lang", "en")]
        wb = load_values(fname, book.get("path", "excel"))
        chap = book["chapter"]
        qs, ans = {}, {}
        for sheet in book["sheets"]:
            ws = wb[sheet["sheet"]]

            def get(a, ws=ws):
                return ws[a].value

            section = sheet.get("section", "structured")
            key = chap if section == "structured" else f"{chap} {section}"
            if sheet["mode"] == "q" and section != "quiz":
                qs.setdefault(key, []).append(render_question(sheet))
            elif sheet["mode"] == "a" and section != "quiz":
                ans.setdefault(key, []).append(render_answers(sheet, get))
            elif sheet["mode"] == "w":
                for b in sheet["blocks"]:
                    if "tag" in b:
                        sections[b["tag"]] = render_answer_block(b, get) if b["k"] != "tb" else r_tb(b)
        for key, items in qs.items():
            sections[f"questions {key}"] = "\n\n".join(items)
        for key, items in ans.items():
            sections[f"answers {key}"] = "\n\n".join(items)
    used = set()
    files = (sorted((ROOT / "notes").glob("*.md")) + sorted((ROOT / "model-questions").glob("*.md"))
             + sorted((ROOT / "prinsip-akaun").rglob("*.md")))
    for path in files:
        text = path.read_text(encoding="utf-8")
        for name, body in sections.items():
            text, hit = replace_section(text, name, body, path)
            if hit:
                used.add(name)
        path.write_text(text, encoding="utf-8")
    unused = sorted(set(sections) - used)
    missing = []
    for path in files:
        for name in re.findall(r"<!-- BEGIN GENERATED: (.*?) -->", path.read_text(encoding="utf-8")):
            if name not in sections:
                missing.append(f"{path.name}: {name}")
    print("filled", len(used), "sections")
    if unused:
        print("not placed:", ", ".join(unused))
    if missing:
        print("MISSING:", "; ".join(missing))
        sys.exit(1)


if __name__ == "__main__":
    main()
